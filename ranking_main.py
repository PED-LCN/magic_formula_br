import pandas as pd
import fundamentus
import yfinance as yf
import numpy as np
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

FILTRO_LIQUIDEZ_MINIMA = 6000000 
FILTRO_RISCO_DECIL = 0.08         

LIMITES_FIXOS = {
    'ey':  {'max': 0.203502579, 'min': 0.035886409},
    'btm': {'max': 1.82251566,  'min': 0.118213687}
}

def limpar_dados(df):
    df.columns = df.columns.str.strip().str.lower()
    mapa = {
        'patrliq': 'patrimonio', 'patrim_liq': 'patrimonio', 
        'liq2m': 'liquidez', 'pebit': 'p_ebit', 'evebit': 'ev_ebit',
        'cotacao': 'cotacao', 'pvp': 'pvp', 'pl': 'pl'
    }
    df = df.rename(columns=mapa)
    cols = ['cotacao', 'ev_ebit', 'liquidez', 'patrimonio', 'pvp', 'pl']
    for col in cols:
        if col not in df.columns: df[col] = 0.0
        else: df[col] = pd.to_numeric(df[col], errors='coerce').fillna(0)
    return df

def get_yahoo_data(tickers):
    print(f"   > Consultando Yahoo Finance para {len(tickers)} ativos...")
    tickers_sa = [f"{t}.SA" for t in tickers]
    
    dict_vol = {}
    
    try:
        
        hist = yf.download(tickers_sa, period="1y", interval="1d", progress=False)['Close']
        
        if isinstance(hist, pd.Series): 
            hist = hist.to_frame()

        if hist.empty:
            return {}
        
        for col in hist.columns:
            t_name = str(col[0]) if isinstance(col, tuple) else str(col)
            ticker_clean = t_name.replace('.SA', '')
            retornos = hist[col].pct_change(fill_method=None).dropna()

            if len(retornos) > 100:
                dict_vol[ticker_clean] = retornos.std() * np.sqrt(252)
            else:
                dict_vol[ticker_clean] = 999.0
                
    except Exception as e:
        print(f"   [Aviso] Erro leve no download de histórico: {e}")

    return dict_vol

def buscar_nomes_top(tickers):
    
    print(f"   > Buscando nomes das empresas para o Top {len(tickers)}...")
    nomes = {}
    for t in tickers:
        try:
            
            ticker_obj = yf.Ticker(f"{t}.SA")
            info = ticker_obj.info
            nomes[t] = info.get('shortName', info.get('longName', t))
        except:
            nomes[t] = t 
    return nomes

def winsorizar_fixo(series, nome_metrica):
    teto = LIMITES_FIXOS[nome_metrica]['max']
    piso = LIMITES_FIXOS[nome_metrica]['min']
    return series.clip(lower=piso, upper=teto)

def main():
    print("1. Obtendo dados do mercado (Fundamentus)...")
    try:
        df = fundamentus.get_resultado()
    except Exception as e:
        print(f"Erro crítico ao acessar Fundamentus: {e}")
        return

    df = limpar_dados(df)

    print("2. Selecionando ticker mais líquido por empresa...")
    df['prefixo'] = df.index.str[:4]
    df = df.sort_values('liquidez', ascending=False)
    df = df.groupby('prefixo', group_keys=False).head(1)
    
    df = df[df['liquidez'] > FILTRO_LIQUIDEZ_MINIMA]
    if 'pl' in df.columns: df = df[df['pl'] > 0] 

    dict_vol = get_yahoo_data(df.index.tolist())
    df['volatilidade'] = df.index.map(dict_vol).fillna(999.0)
    
    corte_vol = df['volatilidade'].quantile(1.0 - FILTRO_RISCO_DECIL)
    df = df[df['volatilidade'] < corte_vol]
    print(f"   > Carteira candidata: {len(df)} ativos")

    df['btm'] = df.apply(lambda x: 1/x['pvp'] if x['pvp'] > 0 else 0, axis=1)
    
    def calc_ey(row):
        if row['ev_ebit'] <= 0: return (1 / row['pl']) if row['pl'] > 0 else 0
        return (1 / row['ev_ebit'])
    df['ey'] = df.apply(calc_ey, axis=1)

    df['ey_w'] = winsorizar_fixo(df['ey'], 'ey')
    df['btm_w'] = winsorizar_fixo(df['btm'], 'btm')
    
    z_ey = (df['ey_w'] - df['ey_w'].mean()) / df['ey_w'].std()
    z_btm = (df['btm_w'] - df['btm_w'].mean()) / df['btm_w'].std()
    
    df['index_final'] = (z_ey + z_btm) / 2
    df = df.sort_values('index_final', ascending=False)

    top_100_tickers = df.head(100).index.tolist()
    dict_nomes = buscar_nomes_top(top_100_tickers)

    df['nome_empresa'] = [dict_nomes.get(t, t) for t in df.index]

    print("3. Gerando Excel Final...")
    wb = Workbook()
    ws = wb.active
    ws.title = "Carteira Deep Value"

    ws['A1'] = "Screening"
    ws['A1'].font = Font(bold=True, size=11)
    ws['B1'] = "Filtro de liquidez:"
    ws['B1'].font = Font(bold=True)
    
    ws['A2'] = datetime.today().strftime('%d/%m/%Y')
    ws['B2'] = f"R$ {FILTRO_LIQUIDEZ_MINIMA:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    
    headers = [
        "Ranking", "Código", "Nome da Empresa", "Idx ", 
        "EY (Yield)", "BtM (Yield)", "Volatilidade", "Liquidez", "Cotação"
    ]
    
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    verde = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    amarelo = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    vermelho = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    row_header = 5

    border_black = Border(
        left=Side(style='medium', color="000000"),
        right=Side(style='medium', color="000000"),
        top=Side(style='thin', color="000000"),
        bottom=Side(style='thin', color="000000")
    )
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=row_header, column=i, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    rank = 1
    for ticker, row in df.head(60).iterrows():
        r = rank + row_header
        
        if rank <= 20: fill = verde
        elif rank <= 30: fill = amarelo
        else: fill = vermelho
        
        dados = [
            rank, ticker, row['nome_empresa'], row['index_final'], 
            row['ey'], row['btm'], row['volatilidade'], 
            row['liquidez'], row['cotacao']
        ]
        
        for c, val in enumerate(dados, 1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.fill = fill
            cell.alignment = Alignment(horizontal='center')
            
            if c in [4, 5, 6, 7]: cell.number_format = '0.0000'
            if c in [8, 9]: cell.number_format = '"R$ "#,##0.00'

            cell.border = border_black
                
        rank += 1
        
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 25
    for col in ['D','E','F','G','H','I']:
        ws.column_dimensions[col].width = 15

    wb.save("Ranking.xlsx")
    print("Sucesso! Arquivo 'Ranking.xlsx' gerado.")

if __name__ == "__main__":
    main()